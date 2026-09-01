"""交易文件解析：保留物理行号，文本不执行，日期和金额不猜测。"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
SHEETS = {'orders': '订单明细', 'refunds': '退款明细', 'costs': '商品成本'}
FIELDS = {
    'orders': {'order_line_id': '订单明细ID', 'order_id': '订单号', 'product_id': '商品编码', 'product_name': '商品名称', 'paid_date': '支付日期', 'status': '订单状态', 'quantity': '数量', 'paid_amount': '实付金额'},
    'refunds': {'refund_id': '退款流水ID', 'order_line_id': '订单明细ID', 'refund_date': '退款日期', 'status': '退款状态', 'amount': '退款金额'},
    'costs': {'product_id': '商品编码', 'effective_from': '成本生效日期', 'unit_cost': '单位成本'},
}
STATUSES = {'已支付': 'paid', '已发货': 'shipped', '已完成': 'completed', '已退款': 'refunded', '待支付': 'pending', '已取消': 'cancelled', '成功': 'success', '处理中': 'processing', '失败': 'failed'}


def issue(row, field, message):
    return {'row': row, 'field': field, 'message': message}


def empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def text(value, required=True):
    if empty(value):
        if required:
            raise ValueError('必填字段不能为空')
        return ''
    if not isinstance(value, str):
        raise ValueError('标识或文本必须为文本格式，避免 Excel 丢失前导零或长编号精度')
    value = value.strip()
    if len(value) > 150 or any(ord(c) < 32 for c in value):
        raise ValueError('文本超过 150 字或包含控制字符')
    return value


def business_date(value, optional=False):
    if empty(value) and optional:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        raise ValueError('使用 YYYY-MM-DD、YYYY/MM/DD 或 Excel 日期单元格；不猜测数字序列日期')
    value = value.strip().replace('/', '-')
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}:\d{2})?', value):
        raise ValueError('日期格式错误，使用 YYYY-MM-DD，可附 HH:mm:ss（上海当地时间）')
    try:
        return datetime.fromisoformat(value).date().isoformat()
    except ValueError:
        raise ValueError('日期不存在或时间无效') from None


def cents(value, optional=False):
    if empty(value) and optional:
        return None
    raw = str(value).strip()
    if not re.fullmatch(r'\d+(?:\.\d{1,2})?', raw):
        raise ValueError('金额必须非负、最多两位小数，单位元；不接受公式、科学计数、千分位或货币符号')
    number = Decimal(raw)
    if number > 100000000:
        raise ValueError('单条金额超过一亿元上限')
    return int(number * 100)


def read_table(content, filename, kind):
    if kind not in FIELDS:
        raise ValueError('类型必须为 orders、refunds 或 costs')
    if not content or len(content) > MAX_BYTES:
        raise ValueError('文件为空或超过 5 MB')
    rows, errors, warnings = [], [], []
    suffix = Path(filename).suffix.lower()
    if suffix == '.csv':
        try:
            try:
                decoded = content.decode('utf-8-sig')
                encoding = 'UTF-8'
            except UnicodeDecodeError:
                decoded = content.decode('gb18030')
                encoding = 'GB18030'
            reader = csv.reader(io.StringIO(decoded, newline=''), strict=True)
            previous_line = 0
            for values in reader:
                number = previous_line + 1
                previous_line = reader.line_num
                rows.append((number, values))
                if len(rows) > MAX_ROWS + 1 or len(values) > 30:
                    raise ValueError('文件超过 5000 数据行或 30 列')
            sheet = 'CSV'
        except (UnicodeError, csv.Error) as exc:
            raise ValueError('无法读取 CSV，请保存为 UTF-8 或 GB18030 的逗号分隔文件') from exc
    elif suffix == '.xlsx':
        encoding, sheet = 'XLSX', SHEETS[kind]
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 500 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
                    raise ValueError('Excel 解压内容超过安全限额')
                if len({e.filename for e in entries}) != len(entries):
                    raise ValueError('Excel 存在重复 ZIP 条目')
                if any('vbaProject' in e.filename or 'externalLinks/' in e.filename for e in entries):
                    raise ValueError('不接受含宏或外部链接的 Excel')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if sheet not in book.sheetnames:
                    if len(book.sheetnames) != 1:
                        raise ValueError(f'多工作表文件须含「{sheet}」，请勿用财报汇总表代替交易明细')
                    sheet = book.sheetnames[0]
                if len(book.sheetnames) > 1:
                    warnings.append(issue(1, '工作表', '本次仅处理「' + sheet + '」；其他工作表：' + '、'.join(n for n in book.sheetnames if n != sheet)))
                ws = book[sheet]
                if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > 30:
                    raise ValueError('工作表超过 5000 数据行或 30 列；请清除多余格式行')
                # 不依赖 XLSX 声明的 dimensions，防止其少报行列导致静默遗漏。
                ws.reset_dimensions()
                for number, cells in enumerate(ws.iter_rows(), 1):
                    if number > MAX_ROWS + 1 or len(cells) > 30:
                        raise ValueError('工作表超过 5000 数据行或 30 列')
                    values = []
                    for i, cell in enumerate(cells, 1):
                        values.append(cell.value)
                        if cell.data_type in ('f', 'e'):
                            errors.append(issue(number, f'第 {i} 列', '不接受公式或错误单元格，请粘贴为值；不计算公式'))
                    rows.append((number, values))
            finally:
                book.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('无法读取 Excel；仅支持未加密的标准 .xlsx 文件') from exc
    else:
        raise ValueError('仅支持 .xlsx 与 .csv；旧版 .xls 请另存为 .xlsx')
    if not rows:
        raise ValueError('文件没有表头或数据')
    return rows, errors, warnings, {'sheet': sheet, 'encoding': encoding}


def parse_upload(content, filename, kind):
    rows, errors, warnings, info = read_table(content, filename, kind)
    fields = FIELDS[kind]
    aliases = {v: k for k, v in fields.items()} | {k: k for k in fields}
    header = list(rows[0][1])
    while header and empty(header[-1]):
        header.pop()
    names = [aliases.get(str(h).strip()) for h in header]
    for i, name in enumerate(names):
        if name is None:
            errors.append(issue(1, str(header[i]), '不支持的列；请按模板改名或移除，不能把财报收入当订单实付'))
        elif names.count(name) > 1:
            errors.append(issue(1, fields[name], '重复表头（包括中英文别名重复）'))
    for name in fields:
        if name not in names:
            errors.append(issue(1, fields[name], '缺少必需表头'))
    parsed = []
    if any(e['row'] == 1 for e in errors):
        return parsed, errors, warnings, info | {'total_rows': len(rows) - 1}
    invalid_cells = {e['row'] for e in errors}
    for number, values in rows[1:]:
        if all(empty(v) for v in values):
            warnings.append(issue(number, '整行', '空白行，明确跳过'))
            continue
        if len(values) > len(header) and any(not empty(v) for v in values[len(header):]):
            errors.append(issue(number, '整行', '存在表头之外的数据列'))
            continue
        values = list(values) + [None] * max(0, len(header) - len(values))
        raw = dict(zip(names, values))
        item = {'_row': number}
        row_errors = []
        for key in fields:
            try:
                value = raw[key]
                if key in ('paid_date', 'refund_date', 'effective_from'):
                    item[key] = business_date(value, optional=key == 'paid_date')
                elif key in ('paid_amount', 'amount', 'unit_cost'):
                    item[key + '_cents'] = cents(value, optional=key == 'unit_cost')
                elif key == 'quantity':
                    if not re.fullmatch(r'[1-9]\d{0,5}(?:\.0+)?', str(value).strip()):
                        raise ValueError('数量必须为 1 至 999999 的整数')
                    item[key] = int(Decimal(str(value)))
                elif key == 'status':
                    status = text(value)
                    status = STATUSES.get(status, status)
                    allowed = ('pending', 'cancelled', 'paid', 'shipped', 'completed', 'refunded') if kind == 'orders' else ('processing', 'success', 'failed')
                    if status not in allowed:
                        raise ValueError('不支持的状态，请使用模板中的状态')
                    item[key] = status
                else:
                    item[key] = text(value)
            except ValueError as exc:
                row_errors.append(issue(number, fields[key], str(exc)))
        if not row_errors and kind == 'orders':
            unpaid = item['status'] in ('pending', 'cancelled')
            if not unpaid and item['paid_date'] is None:
                row_errors.append(issue(number, '支付日期', '有效订单必须提供支付日期'))
            if unpaid and (item['paid_date'] is not None or item['paid_amount_cents'] != 0):
                row_errors.append(issue(number, '订单状态', '待支付/已取消要求支付日期为空且实付为 0；退款订单应保留有效支付状态'))
        if not row_errors and kind == 'refunds' and item['amount_cents'] <= 0:
            row_errors.append(issue(number, '退款金额', '退款金额必须大于零'))
        errors.extend(row_errors)
        if not row_errors and number not in invalid_cells:
            parsed.append(item)
    if len(rows) == 1:
        errors.append(issue(2, '整行', '文件没有数据行'))
    return parsed, errors, warnings, info | {'total_rows': len(rows) - 1}
